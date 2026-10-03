"""One row per outlet visit and own beer product, with GPS-derived administrative names."""
import math
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from app.core.summary_marker import is_summary_name
from app.services.offline_locations import resolve_location, get_boundary_index

HEADERS = ['Date', 'Region', 'Dealer', 'Outlet Name', 'Outlet Type',
           'Phone Number Outlet', 'Latitude', 'Longitude', 'Province', 'District', 'Commune', 'Product_Beer', 'Movement Rate']


# Explicit own-brand beer allowlist; avoid matching energy/water brand names.
BEER_PRODUCTS = frozenset(name.casefold() for name in (
    'CB LITE ORD', 'CBC 4.4 NCP', 'CB Original NCP', 'CB LITE NCP',
    'CB BLACK NCP', 'CB Pint', 'CBL Pint', 'CB SUPEEME Pint',
    'CB SUPREME Pint', 'CB Black Pint',
))


def beer_metrics(row):
    return [metric for metric in (getattr(row, 'product_metrics', None) or [])
            if ' '.join(str(getattr(metric, 'product_name', '')).split()).casefold() in BEER_PRODUCTS]


def coordinates(lat, lon):
    try:
        lat, lon = float(lat), float(lon)
        if not math.isfinite(lat) or not math.isfinite(lon) or not (-90 <= lat <= 90 and -180 <= lon <= 180):
            return None
        if lat == 0 and lon == 0:
            return None
        return lat, lon
    except (TypeError, ValueError):
        return None


def create_detail_export(rows, day, output_dir, resolver=resolve_location):
    output_dir = Path(output_dir); output_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = 'Outlet Detail'; ws.append(HEADERS)
    count = unresolved = 0
    attempted = {}
    excel_row = 1
    for row in rows:
        if is_summary_name(getattr(row, 'outlet_name', None)):
            continue
        metrics = beer_metrics(row)
        if not metrics:
            continue
        pin = coordinates(getattr(row, 'gps_latitude', None), getattr(row, 'gps_longitude', None))
        names = ['', '', '']
        if pin:
            key = pin
            names = attempted.get(key)
            if names is None:
                names = list(resolver(*pin))
                attempted[key] = names
        unresolved += int(not all(names))
        values = [day, getattr(row,'region',None), getattr(row,'dealer',None),
                  getattr(row,'outlet_name',None), getattr(row,'outlet_type',None),
                  str(getattr(row,'phone_number',None) or ''),
                  pin[0] if pin else None, pin[1] if pin else None, *names]
        for metric in metrics:
            product = getattr(metric, 'product_name', '') if metric else ''
            rate = getattr(metric, 'movement_score', None) if metric else None
            if metric is not None and getattr(metric, 'available', None) is False:
                rate = 0
            if rate is not None:
                try:
                    rate = float(rate)
                    if not math.isfinite(rate):
                        rate = None
                except (TypeError, ValueError):
                    rate = None
            ws.append(values + [product, rate])
            excel_row += 1
            for column in range(1, len(HEADERS) + 1):
                cell = ws.cell(excel_row, column)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
            ws.cell(excel_row, 1).number_format = 'yyyy-mm-dd'
            ws.cell(excel_row, 6).number_format = '@'
            ws.cell(excel_row, 13).number_format = '0.##'
        count += 1  # Completion message reports visits, not expanded product rows.
    for c in ws[1]:
        c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='174A73')
    from openpyxl.utils import get_column_letter
    for i,width in enumerate([14,12,12,32,22,24,16,16,24,24,24,32,20],1):
        ws.column_dimensions[get_column_letter(i)].width=width
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    path=output_dir / f'Export_detail_{day}.xlsx'
    wb.save(path)
    return path,count,unresolved


def generate_outlet_detail_export(day):
    get_boundary_index()  # Fail clearly if deployment omitted the bundled map.
    from app.core.config import settings
    from app.kobo.sync import fetch_report_submissions_fast
    rows=fetch_report_submissions_fast(None,day,metadata_only=False)
    if not rows:
        raise ValueError(f'No submissions found for {day}')
    return create_detail_export(rows,day,settings.export_path)
