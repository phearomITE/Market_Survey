"""Weekly report guidance. The cycle is keyed to the report date."""

from datetime import date, datetime

from openpyxl.styles import Alignment, Font, PatternFill


NO_COMPROMISE = (
    "1.Mass Products មិនត្រូវឲ្យខ្វះស្លកក្នុងផ្ទះមួយ(CBL/Wurkz/Exprez/Dazz/Water/Sport PET 500ml/Ize PET 500ml)",
    "2.ផលិតផលហួសការកំណត់",
    "3. ការបញ្ចូលរបាយការណ៍លក់មិនត្រឹមត្រូវ",
    "4. របាយការណ៍លក់ជូនអតិថិជនមិនពិត (ទាំងចំនួនលក់ និងតម្លៃ)។",
)

# Exact text from Market_Improvement_STR3_2026-09-19.xlsx, cells A46:A52.
DONT = (
    "1.កុំធ្វើ Retail នៅកន្លែងដដែលៗ និងកន្លែងលក់ដាច់ (កុំដណ្ដើមមួយWholesale)",
    "2.កុំឡើងផលិតផលដែលលក់ដាច់ស្រាប់ (Mass Product)",
    "3.កុំដើររំលងមួយ ឬទុកទីតាំងចោលយូរ",
    "4.កុំយក Program ទៅ Build ម៉ូយដដែលៗ",
    "5. កុំយកធលិតផលដាក់មួយធំៗច្រើនពេក ផលិតផលលក់យឺត និងផលិតផលថ្មី",
    "6. កុំប្រជុំយូរពេក (Morning Talk កុំឲ្យលើស 15នាទី)",
    "7.កុំសន្យាជាមួយមួយបើមិនច្បាស់លាស់។",
)


def guidance_for_date(report_date):
    """12 Sep 2026 = No Compromise, 19 Sep = Don't; repeat weekly."""
    if isinstance(report_date, datetime):
        report_date = report_date.date()
    elif isinstance(report_date, str):
        report_date = next(
            (datetime.strptime(report_date.strip(), fmt).date()
             for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
             if _matches_date(report_date.strip(), fmt)),
            None,
        )
    if not isinstance(report_date, date):
        report_date = date.today()
    weeks = (report_date - date(2026, 9, 12)).days // 7
    return ("No Compromise", NO_COMPROMISE) if weeks % 2 == 0 else ("Don't", DONT)


def _matches_date(value, fmt):
    try:
        datetime.strptime(value, fmt)
        return True
    except ValueError:
        return False


def write_guidance(ws, report_date, header_row):
    """Put dated guidance in the left A:G block beside issues and suggestions."""
    title, lines = guidance_for_date(report_date)
    header = ws.cell(header_row, 1)
    header.value = title
    header.fill = PatternFill("solid", fgColor="FF1018")
    header.font = Font(name="Noto Sans Khmer", size=12, bold=True, color="FFFFFF")
    header.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[header_row].height = 26
    for index, line in enumerate(lines, 1):
        row = header_row + index
        for first, last in ((1, 7), (9, 17), (19, 27)):
            address = f"{ws.cell(row, first).coordinate}:{ws.cell(row, last).coordinate}"
            if address not in {str(merged) for merged in ws.merged_cells.ranges}:
                ws.merge_cells(start_row=row, start_column=first,
                               end_row=row, end_column=last)
        cell = ws.cell(row, 1)
        cell.value = line if line.lstrip().startswith(f"{index}.") else f"{index}. {line}"
        cell.font = Font(name="Noto Sans Khmer", size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 32 if len(line) < 80 else 50
    return header_row + len(lines)
