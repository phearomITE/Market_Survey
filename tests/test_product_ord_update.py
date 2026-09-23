import ast
from contextlib import closing
import unittest
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


class ProductOrdUpdateTests(unittest.TestCase):
    def test_competitor_field_resolves_for_report_and_raw_export(self):
        """Execute the actual field resolver without importing the DB runtime."""
        tree = ast.parse((ROOT / 'app/reports/aggregator.py').read_text(encoding='utf8'))
        function = next(node for node in tree.body
                        if isinstance(node, ast.FunctionDef) and node.name == 'competitor_field')
        scope = {
            'lru_cache': lru_cache,
            'COMPETITOR_CODES': {'Boostrong ORD': ['boostrong'], 'EXPREZ Can 330ml ORD': ['exprez_can_330']},
            'OWN_PRODUCTS': ['EXPREZ Can 330ml ORD'],
            'slug': lambda product: product.lower().replace(' ', '_'),
            '_field_label_aliases': lambda product, field: [],
            'product_field': lambda product, field: [f'own_{field}_{product}'],
        }
        exec(compile(ast.Module(body=[function], type_ignores=[]), '<aggregator>', 'exec'), scope)
        self.assertIn('comp_mov_boostrong', scope['competitor_field']('Boostrong ORD', 'mov'))
        self.assertIn('own_mov_EXPREZ Can 330ml ORD',
                      scope['competitor_field']('EXPREZ Can 330ml ORD', 'mov'))

    def test_product_lists_and_comparison_groups(self):
        tree = ast.parse((ROOT / 'app/reports/aggregator.py').read_text(encoding='utf8'))
        values = {}
        for node in tree.body:
            if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
                try: values[node.targets[0].id] = ast.literal_eval(node.value)
                except ValueError: pass
        own = values['OWN_PRODUCTS']
        competitor = values['COMPETITOR_PRODUCTS']
        self.assertIn('CAMBODIA ED ORD', own)
        self.assertIn('DAZZ Zero Sugar ORD', own)
        self.assertIn('EXPREZ Can 330ml ORD', own)
        self.assertIn('King Kong Ice ORD', competitor)
        self.assertIn('Idol Can 330ml ORD', competitor)
        self.assertNotIn('CAMBODIA Sport 300mL', own)
        self.assertNotIn('CAMBODIA Sport 300ml', competitor)
        self.assertIn('AIRA', competitor)
        self.assertIn('CAMBODIA COLA', own)
        self.assertFalse(any('300ml' in product.lower() and 'sport' in product.lower()
                             for group in values['OFFTAKE_COMPARE_GROUPS'] for product in group))

    def test_kobo_form_and_export_template(self):
        with closing(load_workbook(ROOT / 'templates/KB_Market_Improvement_XLSForm_GT_HORECA.xlsx', read_only=True)) as form:
            survey = form['survey']
            labels = [cell.value for row in survey for cell in row if isinstance(cell.value, str)]
            self.assertIn('CAMBODIA ED ORD', labels)
            self.assertIn('CAMBODIA Sport 500mL ORD', labels)
            labels = [cell.value for row in survey for cell in row if isinstance(cell.value, str)]
            self.assertFalse(any('CAMBODIA Sport 300mL' in value for value in labels))
        with closing(load_workbook(ROOT / 'templates/template_general.xlsx', read_only=True)) as template:
            self.assertEqual(template.active['B14'].value, 'CAMBODIA ED ORD')


if __name__ == '__main__':
    unittest.main()
