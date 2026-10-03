import tempfile
import unittest
from datetime import date
from types import SimpleNamespace as Row
from openpyxl import load_workbook
from app.services.outlet_detail_export import create_detail_export

class ProductDetailTests(unittest.TestCase):
    def test_product_rows_rates_and_outlet_fields(self):
        metrics = [Row(product_name='CB LITE ORD', available=True, movement_score=10),
                   Row(product_name='CAMBODIA Sport 500mL ORD', available=False, movement_score=7),
                   Row(product_name='Available with missing score', available=True, movement_score=None),
                   Row(product_name='Zero score', available=True, movement_score=0)]
        outlet = Row(outlet_name='Shop',region='R1',dealer='CA2',phone_number='012345678',
                     gps_latitude=11.55,gps_longitude=104.93,product_metrics=metrics)
        with tempfile.TemporaryDirectory() as folder:
            path, visits, unresolved = create_detail_export([outlet],date(2026,1,3),folder,lambda *p:['P','D','C'])
            wb=load_workbook(path);ws=wb.active
            self.assertEqual((visits,unresolved),(1,0))
            self.assertEqual((ws.max_row,ws.max_column),(5,13))
            self.assertEqual([ws.cell(1,c).value for c in (12,13)],['Product','Movement Rate'])
            self.assertEqual([ws.cell(r,13).value for r in range(2,6)],[10,0,None,0])
            self.assertEqual([ws.cell(r,6).value for r in range(2,6)],['012345678']*4)
            self.assertEqual(ws['L3'].value,'CAMBODIA Sport 500mL ORD')
            wb.close()

if __name__ == '__main__':
    unittest.main()
