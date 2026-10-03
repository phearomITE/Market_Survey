"""Run without Kobo, Telegram credentials, database, or a geocoding API key."""
import os
import tempfile
import unittest
from datetime import date
from types import SimpleNamespace, ModuleType
from unittest.mock import patch
from openpyxl import load_workbook
from app.services.offline_locations import BoundaryIndex, get_boundary_index, polygon_contains
from app.services.outlet_detail_export import generate_outlet_detail_export, HEADERS


def square(x, y, size=1):
    return [[x,y],[x+size,y],[x+size,y+size],[x,y+size],[x,y]]


class OfflineDetailTests(unittest.TestCase):
    def test_polygon_holes_and_edges(self):
        polygon = [square(0,0,4), square(1,1)]
        self.assertTrue(polygon_contains(.5,.5,polygon))
        self.assertFalse(polygon_contains(1.5,1.5,polygon))
        self.assertTrue(polygon_contains(0,2,polygon))
        self.assertFalse(polygon_contains(5,5,polygon))

    def test_priority_fallback_and_ambiguity(self):
        records = [dict(code='1', names=['P','D','C'], priority=0, polygons=[[square(0,0)]]),
                   dict(code='2', names=['Old','Old','Old'], priority=1, polygons=[[square(0,0,3)]]),
                   dict(code='3', names=['P','D','Other'], priority=0, polygons=[[square(1,0)]])]
        index = BoundaryIndex(records)
        self.assertEqual(index.resolve(.5,.5), ['P','D','C'])
        self.assertEqual(index.resolve(2,2), ['Old','Old','Old'])
        self.assertEqual(index.resolve(.5,1), ['', '', ''])
        self.assertEqual(index.resolve(10,10), ['', '', ''])
        self.assertEqual(index.resolve(float('nan'),1), ['', '', ''])

    def test_uploaded_boundaries(self):
        index = get_boundary_index()
        self.assertEqual(index.resolve(11.55,104.93), ['Phnom Penh','Chamkar Mon','Tonle Basak'])
        self.assertEqual(index.resolve(11.5564,104.9282), ['Phnom Penh','Boeng Keng Kang','Boeng Keng Kang Ti Muoy'])
        self.assertEqual(index.resolve(0,0), ['', '', ''])
        self.assertEqual(index.resolve(104.93,11.55), ['', '', ''])

    def test_export_without_api_key_and_no_stale_cache(self):
        day = date(2026,1,3)
        config = ModuleType('app.core.config')
        sync = ModuleType('app.kobo.sync')
        rows = [SimpleNamespace(outlet_name='=Literal outlet', phone_number='012345678',
                    gps_latitude=11.55, gps_longitude=104.93, region='R1', dealer='CA2', outlet_type='GT',product_metrics=[SimpleNamespace(product_name='CB LITE ORD',available=True,movement_score=10)]),
                SimpleNamespace(outlet_name='No GPS', phone_number='098765432',product_metrics=[SimpleNamespace(product_name='CB LITE ORD',available=False,movement_score=0)])]
        calls = []
        def fetch(dealer, requested, metadata_only=False):
            calls.append((dealer,requested,metadata_only))
            return rows
        sync.fetch_report_submissions_fast = fetch
        with tempfile.TemporaryDirectory() as folder:
            config.settings = SimpleNamespace(export_path=folder)
            from pathlib import Path
            Path(folder,'outlet_detail_geocode_cache.json').write_text('{"11.550000,104.930000":["WRONG","WRONG","WRONG"]}')
            with patch.dict(os.environ, {'BIGDATACLOUD_API_KEY':''}), patch.dict('sys.modules', {'app.core.config':config,'app.kobo.sync':sync}):
                path, count, unresolved = generate_outlet_detail_export(day)
            self.assertEqual(calls, [(None,day,False)])
            self.assertEqual((count,unresolved),(2,1))
            wb = load_workbook(path); ws = wb.active
            self.assertEqual([c.value for c in ws[1]], HEADERS)
            self.assertEqual([ws.cell(2,c).value for c in (9,10,11)], ['Phnom Penh','Chamkar Mon','Tonle Basak'])
            self.assertEqual(ws['F2'].value,'012345678')
            self.assertEqual(ws['D2'].data_type,'s')
            self.assertIsNone(ws['K3'].value)
            wb.close()

    def test_empty_date_reports_no_submissions(self):
        config, sync = ModuleType('app.core.config'), ModuleType('app.kobo.sync')
        config.settings = SimpleNamespace(export_path='unused')
        sync.fetch_report_submissions_fast = lambda *args, **kwargs: []
        with patch.dict('sys.modules', {'app.core.config':config,'app.kobo.sync':sync}):
            with self.assertRaisesRegex(ValueError,'No submissions found'):
                generate_outlet_detail_export(date(2026,1,3))


if __name__ == '__main__':
    unittest.main()
