import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

from app.core.summary_marker import is_summary_name

ROOT = Path(__file__).resolve().parents[1]


class SummaryOutletVariants(unittest.TestCase):
    def test_control_row_variants_and_regular_outlet(self):
        for text in (
            'សរុបរួម', '"បូកសរុបរួម"', 'សរុបចុងក្រោយ',
            '# ចែ ម៉ៅ , បូកសរុបរួម', 'បូកសរុបរួម,បូកសរុបរួម10',
            'បូកសរុបរួម5ម៉ូយចុងក្រោយ', 'បូក សរុប រួម',
        ):
            with self.subTest(text=text):
                self.assertTrue(is_summary_name(text))
        self.assertFalse(is_summary_name('ចែ ម៉ៅ'))

    def test_xlsform_routing_matches_markers_and_field_removed(self):
        path = ROOT / 'templates/KB_Market_Improvement_XLSForm_GT_HORECA.xlsx'
        with closing(load_workbook(path, read_only=True)) as form:
            survey = form['survey']
            headers = {cell.value: i for i, cell in enumerate(survey[1])}
            rows = {row[headers['name']]: row for row in survey.iter_rows(min_row=2, values_only=True)
                    if row[headers['name']]}
            self.assertNotIn('submitter_name', rows)
            self.assertIn('សរុបចុងក្រោយ', rows['outlet_name'][headers['hint']])
            for name in ('gps_location', 'key_issues_detail', 'initiative_idea_suggestion'):
                cell = 'required' if name.endswith(('detail','suggestion')) else 'relevant'
                expression = rows[name][headers[cell]]
                self.assertIn("contains(translate(normalize-space(${outlet_name}), ' ', ''), 'សរុបចុងក្រោយ')", expression)
                self.assertIn("contains(translate(normalize-space(${outlet_name}), ' ', ''), 'បូកសរុបរួម')", expression)


if __name__ == '__main__':
    unittest.main()
