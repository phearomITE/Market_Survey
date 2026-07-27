from __future__ import annotations

import os
from datetime import date
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")

from openpyxl import load_workbook

from app.data.dealers import REGION_DEALERS
from app.reports.summary_report import create_summary_report


def _rows():
    result = []
    for region, dealers in REGION_DEALERS.items():
        for index, dealer in enumerate(dealers):
            movement = [4, 6, 9, None][index % 4]
            result.append(
                {
                    "region": region,
                    "dealer": dealer,
                    "member": "7",
                    "total_submissions": 10,
                    "total_outlets": 9,
                    "status": "✅",
                    "cb_lite_movement": movement,
                    "movement_lt5": movement if movement is not None and movement < 5 else None,
                    "movement_5_8": movement if movement is not None and 5 <= movement <= 8 else None,
                    "movement_9_10": movement if movement is not None and 9 <= movement <= 10 else None,
                    "product_competitor": "GB SNOW NCP" if movement in (4, 6) else "",
                    "movement_lead": 10 if movement in (4, 6) else None,
                }
            )
    return result


def test_summary_colors_detail_0_to_8_and_valid_excel_table(tmp_path: Path):
    output = tmp_path / "summary.xlsx"
    detail_rows = [
        {
            "date": date(2026, 7, 25),
            "region": "R1",
            "dealer": "CA1",
            "outlet_name": "Outlet A",
            "phone_number": "+855123",
            "outlet_type": "Drink Shop",
            "stock_status": "គ្រប់",
            "freshness_date": "07.2027",
            "movement_0_8": 8,
            "product_competitor": "GB SNOW NCP",
            "movement_lead": 10,
            "link_map": "https://www.google.com/maps?q=11.1,104.9",
        },
        {
            "date": date(2026, 7, 25),
            "region": "R1",
            "dealer": "CA1",
            "outlet_name": "Outlet B",
            "phone_number": "0",
            "outlet_type": "Wholesale",
            "stock_status": "ខ្វះ",
            "freshness_date": "06.2027",
            "movement_0_8": 0,
            "product_competitor": "Hanuman LITE NCP",
            "movement_lead": 10,
            "link_map": "https://www.google.com/maps?q=11.2,104.8",
        },
    ]

    create_summary_report(_rows(), date(2026, 7, 25), detail_rows, output)
    workbook = load_workbook(output)
    summary = workbook["Summary"]
    detail = workbook["Detail"]

    assert summary["F4"].fill.fgColor.rgb[-6:] == "60497A"
    assert summary["G4"].fill.fgColor.rgb[-6:] == "963634"
    assert summary["H4"].fill.fgColor.rgb[-6:] == "00B050"
    assert summary["F8"].fill.fgColor.rgb[-6:] == "1F4E78"
    assert summary["G8"].fill.fgColor.rgb[-6:] == "60497A"
    assert summary["H8"].fill.fgColor.rgb[-6:] == "963634"
    assert summary["I8"].fill.fgColor.rgb[-6:] == "00B050"

    assert detail["I1"].value == "0 to 8"
    assert detail["I2"].value == 8
    assert detail["I3"].value == 0
    assert detail["L2"].hyperlink.target.endswith("11.1,104.9")
    assert list(detail.tables.keys()) == ["DetailMovementTable"]
    assert detail.tables["DetailMovementTable"].ref == "A1:L3"

    with ZipFile(output) as archive:
        table_xml = ET.fromstring(archive.read("xl/tables/table1.xml"))
    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    columns = table_xml.find("x:tableColumns", namespace)
    ids = [int(column.attrib["id"]) for column in columns]
    assert ids == list(range(1, 13))
