import ast
from contextlib import closing
from datetime import date
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook

from app.reports.guidance import guidance_for_date, write_guidance


ROOT = Path(__file__).resolve().parents[1]


class OrdGuidanceTests(unittest.TestCase):
    def test_weekly_cycle_uses_requested_report_date(self):
        self.assertEqual(guidance_for_date(date(2026, 9, 12))[0], 'No Compromise')
        self.assertEqual(guidance_for_date('2026-09-19')[0], "Don't")
        self.assertEqual(guidance_for_date('26/09/2026')[0], 'No Compromise')
        self.assertEqual(guidance_for_date(date(2026, 10, 3))[0], "Don't")

    def test_guidance_is_visible_in_printable_sheet(self):
        sheet = Workbook().active
        end = write_guidance(sheet, date(2026, 9, 19), 45)
        self.assertEqual(sheet['A45'].value, "Don't")
        self.assertEqual(end, 52)
        self.assertEqual(sheet['A52'].value,
                         '7.កុំសន្យាជាមួយមួយបើមិនច្បាស់លាស់។')
        self.assertTrue(any(str(rng) == 'A52:G52' for rng in sheet.merged_cells.ranges))
        self.assertFalse(any(str(rng) == 'A45:AA45' for rng in sheet.merged_cells.ranges))
        source = (ROOT / 'app/reports/excel_report.py').read_text(encoding='utf8')
        self.assertIn('write_guidance(ws, rdate, layout["summary_header"])', source)
        self.assertIn('max(layout[\'print_end\'], guidance_end)', source)

    def test_wording_matches_uploaded_reports(self):
        from app.reports.guidance import DONT, NO_COMPROMISE
        self.assertEqual(DONT[5],
                         '6. កុំប្រជុំយូរពេក (Morning Talk កុំឲ្យលើស 15នាទី)')
        self.assertEqual(NO_COMPROMISE[0],
                         '1.Mass Products មិនត្រូវឲ្យខ្វះស្លកក្នុងផ្ទះមួយ(CBL/Wurkz/Exprez/Dazz/Water/Sport PET 500ml/Ize PET 500ml)')

    def test_form_and_product_lists_have_both_ord_names(self):
        with closing(load_workbook(ROOT / 'templates/KB_Market_Improvement_XLSForm_GT_HORECA.xlsx', read_only=True)) as form:
            labels = [cell.value for sheet in form for row in sheet for cell in row if isinstance(cell.value, str)]
        for name in ('WURKZ ORD', 'Dragon ORD'):
            self.assertIn(name, labels)
        source = (ROOT / 'app/reports/aggregator.py').read_text(encoding='utf8')
        tree = ast.parse(source)
        lists = {node.targets[0].id: ast.literal_eval(node.value)
                 for node in tree.body if isinstance(node, ast.Assign)
                 and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name)
                 and node.targets[0].id in {'OWN_PRODUCTS', 'COMPETITOR_PRODUCTS'}}
        self.assertIn('WURKZ ORD', lists['OWN_PRODUCTS'])
        self.assertIn('Dragon ORD', lists['COMPETITOR_PRODUCTS'])
        self.assertNotIn('WURKZ', lists['OWN_PRODUCTS'])
        self.assertNotIn('Dragon', lists['COMPETITOR_PRODUCTS'])
        # These are the shipped report templates. User repositories can contain
        # older optional files such as template_channel_specialist.xlsx; the
        # active HORECA path uses template_horeca_products.xlsx.
        for name in ('template_general.xlsx', 'template_by_dealer.xlsx',
                     'template_horeca.xlsx', 'template_horeca_products.xlsx',
                     'template_gt_summary.xlsx'):
            template_path = ROOT / 'templates' / name
            with closing(load_workbook(template_path, read_only=True)) as template:
                labels = [cell.value for row in template.active for cell in row if isinstance(cell.value, str)]
            self.assertNotIn('WURKZ', labels, template_path.name)
            self.assertNotIn('Dragon', labels, template_path.name)

    def test_database_migration_never_deletes_canonical_names(self):
        source = (ROOT / 'app/db/database.py').read_text(encoding='utf8')
        self.assertIn('if old_name == new_name:', source)
        self.assertIn('"WURKZ": "WURKZ ORD"', source)
        self.assertIn('"Dragon": "Dragon ORD"', source)
        self.assertNotIn('"DAZZ ORD": "DAZZ ORD"', source)


if __name__ == '__main__':
    unittest.main()
