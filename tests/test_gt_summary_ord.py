from datetime import date
from types import SimpleNamespace as NS
import pytest
from openpyxl import load_workbook
from app.reports.aggregator import COMPETITOR_PRODUCTS, competitor_field
from app.reports.summary_report import _dealer_movement, create_summary_report


def metric(name, score):
    return NS(product_name=name, movement_score=score, stock_status='full', available=True)


def submission(winner='CB LITE ORD'):
    competitors = ['GB SNOW ORD', 'Hanuman LITE ORD', 'Greet LITE ORD']
    return NS(id=1, submission_id=1, report_date=date(2026, 10, 3),
              region='R1', dealer='CA1', member_no=1, outlet_name='Shop',
              outlet_type='Wholesale', phone_number='012345678',
              gps_latitude=11.55, gps_longitude=104.91,
              product_metrics=[metric('CB LITE ORD', 10 if winner == 'CB LITE ORD' else 2),
                               metric('CB LITE NCP', 10)],
              competitor_metrics=[metric(n, 10 if n == winner else 2) for n in competitors]
                  + [metric('GB SNOW NCP', 10), metric('Hanuman LITE NCP', 10),
                     metric('Greet LITE NCP', 10)], ring_pull_metrics=[])


@pytest.mark.parametrize('winner', ['GB SNOW ORD', 'Hanuman LITE ORD', 'Greet LITE ORD'])
def test_summary_compares_ord_and_ignores_conflicting_ncp(winner):
    result = _dealer_movement([submission(winner)], wide_map={})
    assert result['competitor'] == winner
    assert result['competitor_display'] == 10
    assert result['own_display'] < 10


def test_own_ord_wins_and_summary_headers_are_ord(tmp_path):
    row = submission()
    result = _dealer_movement([row], wide_map={})
    assert result['own_display'] == 10
    assert result['competitor'] is None
    rows = [dict(region='R1', dealer='CA1', total_submissions=1,
                 total_outlets=1, target=None, status='✅')]
    output = create_summary_report(rows, row.report_date,
                                   output_path=tmp_path / 'summary.xlsx',
                                   submissions=[row], report_type='GT')
    ws = load_workbook(output, data_only=True).active
    assert [ws.cell(4, col).value for col in (13,14,15)] == [
        'GB SNOW ORD', 'Hanuman LITE ORD', 'Greet LITE ORD']
    assert 'CB LITE ORD' in ws['K6'].value
    assert ws['M8'].value == 10


def test_greet_ord_form_field_mapping_is_distinct_from_ncp():
    assert 'Greet LITE ORD' in COMPETITOR_PRODUCTS
    fields = competitor_field('Greet LITE ORD', 'mov')
    assert any('greet_lite_ord' in key for key in fields)
    assert not any('greet_lite_ncp' in key for key in fields)


def test_horeca_still_uses_pint():
    row = submission()
    row.product_metrics = [metric('CBL Pint', 2)]
    row.competitor_metrics = [metric('Tiger Crystal Pint', 10)]
    result = _dealer_movement([row], report_type='HORECA', wide_map={})
    assert result['competitor'] == 'Tiger Crystal Pint'


def test_mapping_revision_invalidates_old_cache_hashes():
    import hashlib
    import json
    from app.kobo.sync import _source_hash
    raw = {"_id": 123, "comp_movement_score_greet_lite_ord": 10}
    old = hashlib.sha256(json.dumps(raw, ensure_ascii=False, sort_keys=True,
        default=str, separators=(",", ":")).encode("utf-8")).hexdigest()
    assert _source_hash(raw) != old
    assert _source_hash(raw) == _source_hash(dict(reversed(list(raw.items()))))


@pytest.mark.parametrize('winner', ['CB LITE ORD', 'GB SNOW ORD', 'Hanuman LITE ORD', 'Greet LITE ORD'])
def test_fast_summary_fetch_parses_ord_from_kobo(monkeypatch, winner):
    from app.kobo import sync
    from app.kobo.client import KoboClient
    codes = {'CB LITE ORD': 'cb_lite_ord', 'GB SNOW ORD': 'gb_snow_ord',
             'Hanuman LITE ORD': 'hanuman_lite_ord', 'Greet LITE ORD': 'greet_lite_ord'}
    raw = {'_id': 432, 'report_date': '2026-10-03', 'dealer': 'CA1',
           'region': 'R1', 'outlet_name': 'Test shop', 'outlet_type': 'Wholesale',
                      'fresh_movement_score_cb_lite_ord': 10 if winner == 'CB LITE ORD' else 2,
           'fresh_movement_score_cb_lite_ncp': 10}
    for product, code in codes.items():
        if product != 'CB LITE ORD':
            raw['comp_movement_score_' + code] = 10 if winner == product else 2
    monkeypatch.setattr(KoboClient, 'fetch_submissions', lambda *a, **kw: [raw])
    rows = sync._build_report_submissions(dealer=None, report_date=date(2026, 10, 3),
        wanted=set(), summary_only=True, metadata_only=False)
    assert len(rows) == 1
    assert 'CB LITE ORD' in [m.product_name for m in rows[0].product_metrics]
    assert {'GB SNOW ORD', 'HANUMAN LITE ORD', 'Greet LITE ORD'} <= set(m.product_name for m in rows[0].competitor_metrics)
    result = _dealer_movement(rows, wide_map={})
    if winner == 'CB LITE ORD':
        assert result['own_display'] == 10
        assert result['competitor'] is None
    else:
        assert result['competitor'] == winner
        assert result['competitor_display'] == 10
        assert 0 < result['own_display'] < 10


def test_summary_matches_full_dealer_aggregation():
    from app.reports.aggregator import aggregate_submissions
    row = submission()
    row.competitor_metrics.append(metric('Krud LITE ORD', 10))
    full = aggregate_submissions([row], wide_map={}, include_ring_pull=False,
                                include_manual_summary=False)
    assert _dealer_movement([row], wide_map={})['own_display'] == full['products']['CB LITE ORD']['mov']


@pytest.mark.parametrize('name', ['សរុបរួម', 'បូកសរុបរួម10', 'សរុបចុងក្រោយ CA1', 'បូកសរុបរួម5ម៉ូយចុងក្រោយ'])
def test_summary_counts_only_ordinary_outlets(name, tmp_path, monkeypatch):
    from app.reports import summary_report as sr
    ordinary = submission()
    ordinary2 = submission()
    ordinary2.dealer = 'CA8'
    ordinary2.product_metrics[0].available = False
    control = submission()
    control.outlet_name = name
    monkeypatch.setattr(sr, '_coverage_admin_keys', lambda s: ('0101', '010101'))
    records = [ordinary, ordinary2, control]
    rows = sr.build_summary_rows(records)
    assert next(r for r in rows if r['dealer'] == 'CA1')['total_submissions'] == 1
    out = sr.create_summary_report(rows, ordinary.report_date,
        output_path=tmp_path/'coverage.xlsx', submissions=records)
    ws = load_workbook(out, data_only=True).active
    assert ws.max_column == 15
    assert ws['G5'].value == 2
    assert ws['H5'].value == 1
    assert ws['D5'].value == 1  # shared district counted once, not twice
    assert ws['E5'].value == 1
    assert ws['G8'].value == 1
    assert ws['H8'].value == 1
    assert ws['D18'].value == 1
    assert ws['E18'].value == 1
    assert ws['G18'].value == 2
    assert ws['H18'].value == 1


def test_missing_gps_not_counted_as_an_area():
    from app.reports.summary_report import _coverage_admin_keys
    row = submission()
    row.gps_latitude = None
    assert _coverage_admin_keys(row) == (None, None)


def test_final_summary_only_is_not_a_submitted_dealer():
    from app.reports.summary_report import build_summary_rows
    row = submission()
    row.outlet_name = 'សរុបចុងក្រោយ10 CA1'
    result = next(r for r in build_summary_rows([row]) if r['dealer'] == 'CA1')
    assert result['total_submissions'] == 0
    assert result['total_outlets'] == 0
    assert 'No Submit' in result['status']


def test_cb_only_areas_gb_count_missing_sheet_and_villages(tmp_path, monkeypatch):
    from app.reports import summary_report as sr
    present = submission()
    present.village = 'Village One'
    absent = submission()
    absent.outlet_name = 'No beer shop'
    absent.phone_number = '012000111'
    absent.product_metrics[0].available = False
    absent.competitor_metrics[0].movement_score = 0
    absent.village = 'Village Two'
    summary = submission()
    summary.outlet_name = 'សរុបរួម10'
    monkeypatch.setattr(sr, '_coverage_admin_keys', lambda s: ('0101','010101') if s is present else ('0201','020101'))
    output = sr.create_summary_report([], present.report_date, output_path=tmp_path/'test.xlsx', submissions=[present, absent, summary])
    wb = load_workbook(output)
    ws = wb['Summary_beer']
    assert ws['D5'].value == 1
    assert ws['E5'].value == 1
    assert ws['F5'].value == 1
    assert ws['G5'].value == 2
    assert ws['H5'].value == 1
    assert ws['I5'].value == 1
    detail = wb['Location_no_CB_LITE']
    assert detail.max_row == 2
    assert detail['D2'].value == 'No beer shop'
    assert detail['E2'].value == '012000111'
    assert detail['L2'].value == 'No'
    assert detail['K2'].hyperlink.target.startswith('https://www.google.com/maps?q=')
    assert ws['D18'].font.color.rgb == '00FF0000'


def test_unknown_village_is_not_replaced_by_commune():
    from app.reports.summary_report import _coverage_village_key
    assert _coverage_village_key(submission()) is None
