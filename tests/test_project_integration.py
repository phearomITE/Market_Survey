from datetime import date
from types import SimpleNamespace
from unittest.mock import patch, AsyncMock
import asyncio
from fastapi.testclient import TestClient
from fastapi import FastAPI
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from openpyxl import load_workbook
from app.core.config import settings
from app.db.models import Base, KoboSubmission
from app.web import power_bi
from app.services import report_service
from app.bot import handlers


def test_status_command_uploads_workbook(tmp_path):
    rows = [SimpleNamespace(dealer='CA1', outlet_name='បូកសរុបរួម10')]
    message = SimpleNamespace(reply_text=AsyncMock(return_value=SimpleNamespace(edit_text=AsyncMock())), reply_document=AsyncMock())
    with patch.object(report_service, 'fetch_report_submissions_fast', return_value=rows), patch.object(settings, 'export_dir', str(tmp_path)):
        asyncio.run(handlers.export_status_cmd(SimpleNamespace(effective_message=message), SimpleNamespace(args=['2026-09-19'])))
    message.reply_document.assert_awaited_once()
    files = list(tmp_path.glob('Summary_Status_*.xlsx'))
    wb = load_workbook(files[0]); ws=wb.active
    assert ws.max_row == 66
    statuses = {r[2]:r[3] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert statuses['CA1']=='Submitted Summary'
    wb.close()


def test_csv_auth_and_all_feeds(tmp_path):
    engine=create_engine('sqlite:///'+str(tmp_path/'bi.db'))
    Base.metadata.create_all(engine)
    session_factory=sessionmaker(bind=engine)
    app=FastAPI(); app.include_router(power_bi.router)
    client=TestClient(app)
    with patch.object(power_bi, 'SessionLocal', session_factory), patch.object(settings, 'power_bi_api_key', 'test-only'):
        for table in ['outlets','products','competitors','ring_pulls']:
            assert client.get('/api/power-bi/'+table).status_code==401
            response=client.get('/api/power-bi/'+table, params={'api_key':'test-only'})
            assert response.status_code==200
            assert 'report_date' in response.text
    engine.dispose()


def test_all_advertised_commands_registered():
    import os
    from app.bot.run_bot import _build_application
    clean_env = {k:v for k,v in os.environ.items() if not k.lower().endswith('_proxy')}
    with patch.dict(os.environ, clean_env, clear=True), patch.object(settings, 'telegram_bot_token', '123456:TEST_TOKEN'):
        bot = _build_application()
    commands = {c for group in bot.handlers.values() for handler in group for c in handler.commands}
    assert {'start','help','status','sync_kobo','debug_kobo','report','report_multi','report_today','summary','raw_movement','export','export_status','alert_submit','map'} <= commands
