import ast
from pathlib import Path
from openpyxl import Workbook, load_workbook
ROOT = Path(__file__).resolve().parents[1]

def test_legacy_label_and_report_rendering():
    tree = ast.parse((ROOT/'app/reports/excel_report.py').read_text(encoding='utf8'))
    selected = [n for n in tree.body if (isinstance(n, ast.Assign) and any(isinstance(t,ast.Name) and t.id=='PRODUCT_NAME_MAP' for t in n.targets)) or (isinstance(n,ast.FunctionDef) and n.name in {'_product_key','_force_rewrite_competitor_blocks'})]
    scope = {'Worksheet': object, '_clean': lambda x: ' '.join(str(x or '').split()), '_layout_rows': lambda *a: {'movement_start':35,'movement_end':36}, '_lookup_competitor_metrics':lambda agg,name:agg['competitors'].get(name), '_write_metrics_cells':lambda ws,r,c,m:setattr(ws.cell(r,c+1),'value',m['mov'])}
    exec(compile(ast.Module(body=selected,type_ignores=[]),'<actual code>','exec'),scope)
    ws=Workbook().active;ws['H35']='BACCHUSE ORD';ws['H36']='BACCHUSE Sugar Free ORD'
    scope['_force_rewrite_competitor_blocks'](ws,{'competitors':{'BACCHUSE':{'mov':7},'BACCHUSE Sugar Free':{'mov':4}}})
    assert (ws['H35'].value,ws['I35'].value)==('BACCHUSE',7)
    assert (ws['H36'].value,ws['I36'].value)==('BACCHUSE Sugar Free',4)
    assert scope['_product_key']('DAZZ ORD')=='DAZZ ORD'

def test_legacy_metric_names():
    tree=ast.parse((ROOT/'app/reports/aggregator.py').read_text(encoding='utf8'))
    f=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='_canonical_product_name')
    scope={'Any':object,'ALL_OWN_PRODUCTS':[],'ALL_COMPETITOR_PRODUCTS':[]}
    exec(compile(ast.Module(body=[f],type_ignores=[]),'<actual code>','exec'),scope)
    assert scope['_canonical_product_name']('BACCHUSE ORD')=='BACCHUSE'
    assert scope['_canonical_product_name']('BACCHUSE Sugar Free ORD')=='BACCHUSE Sugar Free'

def test_uploaded_template_labels():
    w=load_workbook(ROOT/'templates/template_general.xlsx')
    assert w.active['H35'].value=='BACCHUSE'
    assert w.active['H36'].value=='BACCHUSE Sugar Free'
    w.close()
