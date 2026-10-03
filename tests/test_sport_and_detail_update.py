import ast
import tempfile
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from openpyxl import load_workbook
from app.services.outlet_detail_export import create_detail_export, coordinates


def test_sport_group_normalized_and_zero_excluded():
    source=Path(__file__).resolve().parents[1]/'app/reports/aggregator.py'
    tree=ast.parse(source.read_text(encoding="utf-8"))
    names={'_apply_offtake_comparison_goal'}
    nodes=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in names]
    groups=next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='OFFTAKE_COMPARE_GROUPS' for t in n.targets))
    scope={'Any':object,'OFFTAKE_COMPARE_GROUPS':groups,'to_float':lambda x:float(x) if x is not None else None,'to_int':lambda x:int(x or 0),'round_half_up':lambda x:int(x+0.5),'_get_movement_bucket':lambda r,p: ('products',r['products'][p]) if p in r['products'] else None}
    exec(compile(ast.Module(body=nodes,type_ignores=[]),'actual','exec'),scope)
    labels=['CAMBODIA Sport 500mL ORD','Pocari Sweat','V-Active Sport']
    result={'products':{p:{'_mov_effective':v,'_mov_avg':v,'_movement_count':30,'availability':{'Drink Shop':5}} for p,v in zip(labels,[7,6,2])}}
    scope['_apply_offtake_comparison_goal'](result)
    assert [result['products'][p]['mov'] for p in labels]==[10,9,3]
    for data in result['products'].values():data['availability']={}
    scope['_apply_offtake_comparison_goal'](result)
    assert all(d['mov']==0 for d in result['products'].values())


def test_detail_columns_coordinates_cache_and_literals():
    rows=[SimpleNamespace(outlet_name='=2+2',phone_number='012345678',gps_latitude=11.5,gps_longitude=104.9,region='R1',dealer='CA2',outlet_type='Drink Shop')]*2
    calls=[]
    def resolve(*pin):calls.append(pin);return ['Province','District','Commune']
    with tempfile.TemporaryDirectory() as folder:
        path,n,bad=create_detail_export(rows,date(2026,1,3),folder,resolve)
        assert (n,bad,len(calls))==(2,0,1)
        wb=load_workbook(path);ws=wb.active
        assert ws.max_column==13 and ws.max_row==3
        assert ws['D2'].value=='=2+2' and ws['D2'].data_type=='s'
        assert ws['F2'].value=='012345678'
        assert ws['K2'].value=='Commune';wb.close()
        create_detail_export(rows,date(2026,1,3),folder,resolve)
        assert len(calls)==2
    assert coordinates('nan',104) is None
    assert coordinates(91,104) is None

if __name__=='__main__':
    test_sport_group_normalized_and_zero_excluded()
    test_detail_columns_coordinates_cache_and_literals()
    print('Sport comparison and outlet-detail checks passed')
